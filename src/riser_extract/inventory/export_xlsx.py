"""Excel export: one tab per floor + Master + Summary."""

from __future__ import annotations

import json
import sqlite3
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
BAND_FILL = PatternFill("solid", fgColor="FFF2CC")


def export_xlsx(conn: sqlite3.Connection, project_name: str, out_path: Path) -> Path:
    proj = conn.execute(
        "SELECT * FROM projects WHERE name = ?", (project_name,)
    ).fetchone()
    if not proj:
        raise ValueError(f"unknown project: {project_name}")
    pid = proj["id"]

    wb = Workbook()
    # Summary first
    ws_sum = wb.active
    ws_sum.title = "Summary"
    _write_summary(conn, pid, ws_sum, project_name)

    ws_master = wb.create_sheet("Master")
    _write_master(conn, pid, ws_master, project_name)

    docs = conn.execute(
        "SELECT * FROM documents WHERE project_id = ? ORDER BY floor_tab_name",
        (pid,),
    ).fetchall()
    for d in docs:
        name = d["floor_tab_name"] or d["drawing_number"] or f"Sheet{d['id']}"
        # Excel sheet name max 31 chars
        safe = name[:31]
        ws = wb.create_sheet(safe)
        _write_floor(conn, d, ws, project_name)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _write_summary(conn, pid, ws, project_name: str) -> None:
    ws["A1"] = f"{project_name} — HEATING & COOLING TAKE-OFF SUMMARY"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Generated from getriser.sqlite inventory_lines (plan metres / counts)."
    headers = [
        "category",
        "item_id",
        "description",
        "unit",
        "qty_building",
        "where_found",
        "scope",
        "catalog_sku",
        "catalog_status",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(7, col, h)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)

    row = 8
    for line in conn.execute(
        "SELECT * FROM inventory_lines WHERE project_id = ? ORDER BY category, item_id",
        (pid,),
    ):
        ws.cell(row, 1, line["category"])
        ws.cell(row, 2, line["item_id"])
        ws.cell(row, 3, line["description"])
        ws.cell(row, 4, line["unit"])
        ws.cell(row, 5, line["qty_building"])
        ws.cell(row, 6, line["where_found"])
        ws.cell(row, 7, line["scope"])
        ws.cell(row, 8, line["catalog_sku"])
        ws.cell(row, 9, line["catalog_status"])
        row += 1
    ws.auto_filter.ref = f"A7:I{max(7, row - 1)}"


def _write_master(conn, pid, ws, project_name: str) -> None:
    ws["A1"] = f"{project_name} — MASTER ITEM REGISTER"
    ws["A1"].font = Font(bold=True, size=14)
    headers = [
        "category",
        "item_id",
        "description",
        "unit",
        "scope",
        "qty_building",
        "where_found",
        "catalog_sku",
        "catalog_status",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(7, col, h)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
    row = 8
    for line in conn.execute(
        "SELECT * FROM inventory_lines WHERE project_id = ? ORDER BY category, item_id",
        (pid,),
    ):
        for col, key in enumerate(
            [
                "category",
                "item_id",
                "description",
                "unit",
                "scope",
                "qty_building",
                "where_found",
                "catalog_sku",
                "catalog_status",
            ],
            1,
        ):
            ws.cell(row, col, line[key])
        row += 1
    ws.auto_filter.ref = f"A7:I{max(7, row - 1)}"


def _write_floor(conn, doc, ws, project_name: str) -> None:
    tb = json.loads(doc["title_block_json"]) if doc["title_block_json"] else {}
    ws["A1"] = f"{project_name} — {doc['floor_tab_name']} INVENTORY"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Source drawing: {doc['drawing_number']}  |  Scale: {doc['scale'] or tb.get('scale')}"
    ws["A3"] = f"Revision: {doc['revision'] or tb.get('revision')}  |  PDF: {doc['source_pdf']}"
    ws["A4"] = "Quantities are plan counts / plan metres from Phase 1–2 extract (not tender-converted)."

    headers = [
        "line",
        "category",
        "item_id",
        "annotation",
        "linked_to",
        "count",
        "unit",
        "size",
        "measurement_type",
        "scope",
        "symbol_class",
        "confidence",
        "notes",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(7, col, h)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)

    occs = conn.execute(
        "SELECT * FROM occurrences WHERE document_id = ? ORDER BY item_id, id",
        (doc["id"],),
    ).fetchall()

    # Group by category for band rows
    by_cat: dict[str, list] = defaultdict(list)
    for o in occs:
        # category from catalog
        cat_row = None
        if o["item_id"]:
            cat_row = conn.execute(
                "SELECT category FROM catalog_items WHERE item_id = ?",
                (o["item_id"],),
            ).fetchone()
        cat = (cat_row["category"] if cat_row else None) or "Unclassified"
        by_cat[cat].append(o)

    row = 8
    line_i = 1
    for cat in sorted(by_cat.keys()):
        band = ws.cell(row, 1, f"— {cat}")
        band.fill = BAND_FILL
        band.font = Font(bold=True)
        for c in range(2, len(headers) + 1):
            ws.cell(row, c).fill = BAND_FILL
        row += 1
        for o in by_cat[cat]:
            ws.cell(row, 1, f"A{line_i:02d}")
            ws.cell(row, 2, cat)
            ws.cell(row, 3, o["item_id"])
            ws.cell(row, 4, o["annotation"])
            ws.cell(row, 5, o["linked_to"])
            ws.cell(row, 6, o["length_m"] if o["unit"] == "m" and o["length_m"] is not None else o["qty"])
            ws.cell(row, 7, o["unit"])
            ws.cell(row, 8, o["size"])
            ws.cell(row, 9, o["measurement_type"])
            ws.cell(row, 10, o["scope"])
            ws.cell(row, 11, o["symbol_class"])
            ws.cell(row, 12, o["confidence"])
            ws.cell(row, 13, o["notes"])
            line_i += 1
            row += 1

    ws.auto_filter.ref = f"A7:M{max(7, row - 1)}"
