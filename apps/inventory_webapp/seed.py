"""Seed SQLite from the Hubby inventory Excel (or fallback to floors.json)."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DB_PATH = ROOT / "apps" / "inventory_webapp" / "inventory.db"
EXCEL_PATH = Path("/workspace/attachments/25280-SEN Heating and Cooling Inventory - All Floors - last.xlsx")
FLOORS_JSON = ROOT / "evals" / "datasets" / "hubby-premier-inn-london" / "expected" / "floors.json"


SCHEMA = """
CREATE TABLE IF NOT EXISTS projects (
  id INTEGER PRIMARY KEY,
  name TEXT NOT NULL,
  code TEXT,
  notes TEXT
);

CREATE TABLE IF NOT EXISTS drawings (
  id INTEGER PRIMARY KEY,
  project_id INTEGER REFERENCES projects(id),
  floor_tab_name TEXT,
  drawing_number TEXT,
  title TEXT,
  scale TEXT,
  source_pdf TEXT,
  overlay_pdf TEXT
);

CREATE TABLE IF NOT EXISTS inventory_items (
  id INTEGER PRIMARY KEY,
  drawing_id INTEGER REFERENCES drawings(id),
  line TEXT,
  category TEXT,
  annotation TEXT,
  linked_to TEXT,
  count REAL,
  unit TEXT,
  size TEXT,
  tender_estimate REAL,
  measurement_type TEXT,
  scope TEXT,
  notes TEXT,
  status TEXT DEFAULT 'pending',
  amended_count REAL,
  amended_notes TEXT,
  reviewed_at TEXT
);
"""


def seed(conn: sqlite3.Connection) -> None:
    cur = conn.cursor()
    cur.executescript(SCHEMA)

    # Project
    cur.execute(
        "INSERT OR IGNORE INTO projects (id, name, code, notes) VALUES (?, ?, ?, ?)",
        (1, "Hub by Premier Inn Putney", "25280", "Heating & Cooling take-off. First Floor interpolated."),
    )

    if EXCEL_PATH.exists():
        _seed_from_excel(cur)
    elif FLOORS_JSON.exists():
        _seed_from_json(cur)
    else:
        raise FileNotFoundError("No Excel or floors.json found to seed from")

    conn.commit()


def _seed_from_excel(cur: sqlite3.Cursor) -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    drawing_id = 1
    item_id = 1
    for sheet_name in wb.sheetnames:
        if sheet_name == "Summary":
            continue
        ws = wb[sheet_name]
        # crude parse: look for rows with a line code in col A and annotation in col B
        drawing_number = None
        scale = "1:50 @ A1"
        for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
            text = " ".join(str(c or "") for c in row)
            if "DR-M-" in text or "SEN-" in text:
                drawing_number = text.split("—")[0].strip() if "—" in text else text[:60]
            if "1:50" in text:
                scale = "1:50 @ A1"

        cur.execute(
            """INSERT OR IGNORE INTO drawings
               (id, project_id, floor_tab_name, drawing_number, title, scale)
               VALUES (?, 1, ?, ?, ?, ?)""",
            (drawing_id, sheet_name, drawing_number or sheet_name, sheet_name, scale),
        )

        for row in ws.iter_rows(min_row=10, values_only=True):
            if not row or not row[0]:
                continue
            line = str(row[0]).strip()
            if not line or line.startswith("ASSUMPTIONS") or line.startswith("Item"):
                continue
            annotation = str(row[1] or "")[:200]
            linked = str(row[2] or "")[:300]
            try:
                count = float(row[3]) if row[3] is not None else None
            except (TypeError, ValueError):
                count = None
            unit = str(row[4] or "")
            size = str(row[5] or "")
            try:
                tender = float(row[6]) if row[6] is not None else count
            except (TypeError, ValueError):
                tender = count
            notes = str(row[8] or "")[:500] if len(row) > 8 else ""

            measurement = "printed"
            if "ESTIMATE" in notes.upper() or "inferred" in notes.lower() or "amber" in notes.lower():
                measurement = "scaled"
            if "not drawn" in notes.lower() or "NOT DRAWN" in notes:
                measurement = "specified-not-drawn"

            cur.execute(
                """INSERT OR IGNORE INTO inventory_items
                   (id, drawing_id, line, annotation, linked_to, count, unit, size,
                    tender_estimate, measurement_type, scope, notes, status)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'in', ?, 'pending')""",
                (
                    item_id,
                    drawing_id,
                    line,
                    annotation,
                    linked,
                    count,
                    unit,
                    size,
                    tender,
                    measurement,
                    notes,
                ),
            )
            item_id += 1
        drawing_id += 1


def _seed_from_json(cur: sqlite3.Cursor) -> None:
    data = json.loads(FLOORS_JSON.read_text())
    drawing_id = 1
    item_id = 1
    for floor_name, floor in data.get("floors", {}).items():
        cur.execute(
            """INSERT OR IGNORE INTO drawings
               (id, project_id, floor_tab_name, drawing_number, title, scale)
               VALUES (?, 1, ?, ?, ?, ?)""",
            (
                drawing_id,
                floor.get("floor_tab_name", floor_name),
                floor.get("drawing_number"),
                floor_name,
                floor.get("scale", "1:50"),
            ),
        )
        for hkey, h in floor.get("headlines", {}).items():
            cur.execute(
                """INSERT OR IGNORE INTO inventory_items
                   (id, drawing_id, line, annotation, count, unit, tender_estimate,
                    measurement_type, scope, notes, status)
                   VALUES (?, ?, ?, ?, ?, 'no.', ?, 'printed', 'in', ?, 'pending')""",
                (
                    item_id,
                    drawing_id,
                    hkey,
                    h.get("annotation", ""),
                    h.get("count"),
                    h.get("tender"),
                    "Headline from gold set",
                ),
            )
            item_id += 1
        drawing_id += 1


if __name__ == "__main__":
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    if DB_PATH.exists():
        DB_PATH.unlink()
    conn = sqlite3.connect(DB_PATH)
    seed(conn)
    conn.close()
    print(f"Seeded {DB_PATH}")
